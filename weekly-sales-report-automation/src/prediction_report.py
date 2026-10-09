"""CLI and Excel report for scoring movie predictions against actual results.

Usage:
    python -m src.prediction_report --predictions <link or file> --actuals <link or file>
"""

import argparse
import logging
import sys
from contextlib import closing
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from rich.console import Console
from rich.markup import escape
from rich.table import Table

from .excel_report import (
    BAD_FONT,
    GOOD_FONT,
    GRID_BORDER,
    _add_table_polish,
    _write_dataframe,
    _write_findings,
    _write_header_row,
    _write_title,
)
from .history import HISTORY_DB, fill_from_history, open_history, save_actuals, save_scores, write_history
from .main import OUTPUT_DIR, describe_error, next_free_path, setup_logging
from .predictions import (
    Comparison,
    accuracy_stats,
    compare,
    load_actuals,
    load_source,
    parse_title_map,
    read_predictions,
)

FORECASTS = {"predicted": "Predictions", "projection": "Tracker Projection"}
GOOD_ERROR_PCT = 25  # errors within this are shown in green, beyond it in red
MONEY_FORMAT = '"$"#,##0'
PCT_FORMAT = "+0.0%;-0.0%"

console = Console()


def short_money(value: float) -> str:
    if abs(value) >= 1e9:
        return f"${value / 1e9:.2f}B"
    if abs(value) >= 1e6:
        return f"${value / 1e6:.1f}M"
    return f"${value:,.0f}"


def _call(row: pd.Series, forecast: str) -> str:
    return (
        f"{row['title']}, predicted {short_money(row[forecast])} vs "
        f"{short_money(row['actual'])} actual ({row[f'{forecast}_pct_error']:+,.1f}%)"
    )


def prediction_findings(result: Comparison) -> list[str]:
    """Plain-English findings from fixed rules, so the same sheets always give the same text."""
    movies = result.movies
    released = movies[movies["released"]]
    pending = len(movies) - len(released)
    findings = [
        f"{len(released)} of the {len(movies)} matched movies have been released and scored so far"
        + (f"; {pending} {'is' if pending == 1 else 'are'} still waiting on results." if pending else ".")
    ]

    stats = accuracy_stats(movies, "predicted")
    if stats is None:
        return findings + _data_notes(result)

    n = stats["scored"]
    typical = f"The predictions were typically off by {stats['median_abs_pct']:.0f}% (median"
    if stats["mean_abs_pct"] - stats["median_abs_pct"] >= 10:
        typical += f"; the average is {stats['mean_abs_pct']:.0f}% because of a few big misses"
    findings.append(
        typical + f"). {stats['within_25']} of {n} landed within 25% of the actual result, "
        f"and {stats['within_10']} within 10%."
    )

    if stats["over"] != stats["under"]:
        lean = "high" if stats["over"] > stats["under"] else "low"
        findings.append(
            f"They ran {lean} more often than not: {stats['over']} too high vs {stats['under']} too low "
            f"(median error {stats['median_pct']:+.0f}%)."
        )
    total = (
        f"Added up, the predictions came to {short_money(stats['total_forecast'])} against "
        f"{short_money(stats['total_actual'])} actual ({stats['total_pct']:+.1f}%)"
    )
    if abs(stats["total_pct"]) < stats["median_abs_pct"] / 2:
        total += ", so misses in both directions largely cancelled out"
    findings.append(total + ".")

    errors = released["predicted_pct_error"].dropna().abs()
    findings.append(f"Best call: {_call(released.loc[errors.idxmin()], 'predicted')}.")
    findings.append(f"Biggest miss: {_call(released.loc[errors.idxmax()], 'predicted')}.")

    tracker = accuracy_stats(movies, "projection")
    if tracker is not None:
        both = released.dropna(subset=["predicted_pct_error", "projection_pct_error"])
        wins = int((both["predicted_pct_error"].abs() < both["projection_pct_error"].abs()).sum())
        findings.append(
            f"Against the tracker's own projections (typically off by {tracker['median_abs_pct']:.0f}%), "
            f"the predictions were closer on {wins} of {len(both)} movies."
        )
    return findings + _data_notes(result)


def _data_notes(result: Comparison) -> list[str]:
    """Notes about the input data: results restored from history, and titles that couldn't be matched."""
    findings = []
    restored = result.restored_from_history
    if restored:
        findings.append(
            f"{len(restored)} {'movie' if len(restored) == 1 else 'movies'} had no result in the results sheet "
            f"this time, so the last known result was used: {', '.join(restored)}."
        )
    unmatched_actuals, unmatched_predictions = len(result.unmatched_actuals), len(result.unmatched_predictions)
    if unmatched_actuals:
        names = ", ".join(result.unmatched_actuals[:5])
        if unmatched_actuals > 5:
            names += f" and {unmatched_actuals - 5} more"
        if unmatched_actuals == 1:
            findings.append(
                f"1 movie in the results sheet has no matching prediction "
                f"(it may be listed under a different title): {names}."
            )
        else:
            findings.append(
                f"{unmatched_actuals} movies in the results sheet have no matching prediction "
                f"(they may be listed under a different title): {names}."
            )
    if unmatched_predictions:
        findings.append(
            f"{unmatched_predictions} predicted {'movie doesn' if unmatched_predictions == 1 else 'movies don'}'t "
            "appear in the results sheet; see the Unmatched sheet."
        )
    return findings


# (name, description, upper limit of the absolute % error). Each scored movie lands in exactly one band.
ACCURACY_BANDS = [
    ("Nailed it", "within 10%", 10),
    ("Close", "10-25% off", 25),
    ("Off", "25-50% off", 50),
    ("Way off", "more than 50% off", float("inf")),
]


def accuracy_band(abs_pct: float) -> str:
    return next(name for name, _, limit in ACCURACY_BANDS if abs_pct <= limit)


def _movie_label(title: str, pct: float) -> str:
    return f"{title} ({pct:+.1f}%)" if abs(pct) < 10 else f"{title} ({pct:+,.0f}%)"


@dataclass
class ScorecardRow:
    measure: str
    result: str
    movies: list[str] = field(default_factory=list)


def scorecard(result: Comparison) -> list[ScorecardRow]:
    """Simple measures of the predictions' accuracy, each listing the movies behind it."""
    stats = accuracy_stats(result.movies, "predicted")
    if stats is None:
        return []
    scored = result.movies[result.movies["predicted_pct_error"].notna()].copy()
    scored["abs_error"] = scored["predicted_pct_error"].abs()
    n = len(scored)

    def labels(movies: pd.DataFrame) -> list[str]:
        return [_movie_label(t, p) for t, p in zip(movies["title"], movies["predicted_pct_error"])]

    rows = [
        ScorecardRow("Movies scored", f"{n} of {len(result.movies)} matched movies"),
        ScorecardRow("Typical miss (median)", f"{stats['median_abs_pct']:.0f}%"),
    ]
    scored["band"] = scored["abs_error"].map(accuracy_band)
    for name, description, _ in ACCURACY_BANDS:
        in_band = scored[scored["band"] == name].sort_values("abs_error")
        rows.append(ScorecardRow(f"{name} ({description})", f"{len(in_band)} of {n}", labels(in_band)))

    too_high = scored[scored["predicted_pct_error"] > 0].sort_values("predicted_pct_error", ascending=False)
    too_low = scored[scored["predicted_pct_error"] < 0].sort_values("predicted_pct_error")
    rows.append(ScorecardRow("Predicted too high", f"{len(too_high)} of {n}", labels(too_high)))
    rows.append(ScorecardRow("Predicted too low", f"{len(too_low)} of {n}", labels(too_low)))

    tracker = accuracy_stats(result.movies, "projection")
    if tracker is not None:
        both = scored[scored["projection_pct_error"].notna()]
        closer = both[both["abs_error"] < both["projection_pct_error"].abs()].sort_values("abs_error")
        rows.append(ScorecardRow("Closer than the tracker", f"{len(closer)} of {len(both)}", labels(closer)))
        rows.append(ScorecardRow("Tracker's typical miss (median)", f"{tracker['median_abs_pct']:.0f}%"))

    rows.append(
        ScorecardRow(
            "Total predicted vs actual",
            f"{short_money(stats['total_forecast'])} vs {short_money(stats['total_actual'])} ({stats['total_pct']:+.1f}%)",
        )
    )
    return rows


def detailed_stats_rows(result: Comparison) -> tuple[list[tuple[str, list, str]], list[str]]:
    """Returns ((label, [value per forecast], kind) rows, forecast column headers): the full
    side-by-side stats for every forecaster."""
    stats = {f: accuracy_stats(result.movies, f) for f in FORECASTS}
    stats = {f: s for f, s in stats.items() if s is not None}

    def row(label: str, key: str, kind: str):
        return label, [stats[f][key] for f in stats], kind

    rows = [
        row("Movies scored", "scored", "count"),
        row("Median error (either direction)", "median_abs_pct", "pct_abs"),
        row("Average error (either direction)", "mean_abs_pct", "pct_abs"),
        row("Median lean (+ = too high)", "median_pct", "pct"),
        row("Too high", "over", "count"),
        row("Too low", "under", "count"),
        row("Within 10% of actual", "within_10", "count"),
        row("Within 25% of actual", "within_25", "count"),
        row("Total forecast", "total_forecast", "money"),
        row("Total actual", "total_actual", "money"),
        row("Total difference", "total_pct", "pct"),
    ]
    return [(label, values, kind) for label, values, kind in rows], [FORECASTS[f] for f in stats]


MOVIES_COLUMN_WIDTH = 95


def _build_scorecard_sheet(ws, result: Comparison, findings: list[str], as_of: date) -> None:
    _write_title(ws, f"Prediction Accuracy as of {as_of}", span_cols=3)
    rows = scorecard(result)
    if not rows:
        ws.cell(row=3, column=1, value="No released movies to score yet.")
        _write_findings(ws, findings, start_row=5)
        return

    header_row = 3
    _write_header_row(ws, ["Measure", "Result", "Movies"], header_row)
    for offset, row in enumerate(rows, start=1):
        r = header_row + offset
        movies = ", ".join(row.movies)
        for col, value in enumerate([row.measure, row.result, movies], start=1):
            cell = ws.cell(row=r, column=col, value=value or None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = GRID_BORDER
        ws.cell(row=r, column=1).font = Font(bold=True)
        # Excel doesn't auto-fit wrapped rows written this way, so size each from its movie list.
        lines = max(1, -(-len(movies) // (MOVIES_COLUMN_WIDTH - 5)))
        ws.row_dimensions[r].height = 15 * lines + 2
    last_row = header_row + len(rows)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = MOVIES_COLUMN_WIDTH
    ws.freeze_panes = f"A{header_row + 1}"

    last_row = _write_findings(ws, findings, start_row=last_row + 2)
    _write_detailed_stats(ws, result, start_row=last_row + 2)


def _write_detailed_stats(ws, result: Comparison, start_row: int) -> None:
    rows, headers = detailed_stats_rows(result)
    ws.cell(row=start_row, column=1, value="Detailed stats").font = Font(size=12, bold=True, color="1F4E78")
    table = pd.DataFrame([[label, *values] for label, values, _ in rows], columns=["Metric", *headers])
    header_row = start_row + 1
    _write_dataframe(ws, table, header_row)
    formats = {"count": "#,##0", "money": MONEY_FORMAT, "pct": PCT_FORMAT, "pct_abs": "0.0%"}
    for offset, (_, values, kind) in enumerate(rows, start=1):
        for col in range(2, 2 + len(values)):
            cell = ws.cell(row=header_row + offset, column=col)
            if kind.startswith("pct"):
                cell.value = cell.value / 100
            cell.number_format = formats[kind]
    for col, width in {"A": 32, "B": 26, "C": MOVIES_COLUMN_WIDTH}.items():
        ws.column_dimensions[col].width = width


def _pct_cells(ws, header_row: int, last_row: int, columns: list[int]) -> None:
    for col in columns:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=col, max_col=col):
            if isinstance(cell.value, (int, float)):
                cell.font = GOOD_FONT if abs(cell.value) <= GOOD_ERROR_PCT else BAD_FONT
                cell.value = cell.value / 100
                cell.number_format = PCT_FORMAT


def _for_excel(df: pd.DataFrame) -> pd.DataFrame:
    return df.astype(object).where(df.notna(), None)


def _build_movies_sheet(ws, result: Comparison) -> None:
    released = result.movies[result.movies["released"]].copy()

    def closer(row):
        p, t = row["predicted_pct_error"], row["projection_pct_error"]
        if pd.isna(p) or pd.isna(t):
            return None
        return "Predictions" if abs(p) < abs(t) else ("Tracker" if abs(t) < abs(p) else "Tie")

    table = pd.DataFrame(
        {
            "Movie": released["title"],
            "Release Date": released["release_date"],
            "Predicted": released["predicted"],
            "Tracker Projection": released["projection"],
            "Actual": released["actual"],
            "Predicted Error": released["predicted_pct_error"],
            "Tracker Error": released["projection_pct_error"],
            "Closer": released.apply(closer, axis=1) if len(released) else [],
        }
    )
    _write_title(ws, "Movie by Movie: released movies, forecasts vs actual worldwide gross", span_cols=8)
    header_row = 3
    last_row = _write_dataframe(
        ws, _for_excel(table), header_row, number_formats={1: "yyyy-mm-dd", 2: MONEY_FORMAT, 3: MONEY_FORMAT, 4: MONEY_FORMAT}
    )
    _pct_cells(ws, header_row, last_row, columns=[6, 7])
    ws.column_dimensions["A"].width = 38
    _add_table_polish(ws, header_row, max(last_row, header_row), last_col=len(table.columns))


def _build_upcoming_sheet(ws, result: Comparison) -> None:
    upcoming = result.movies[~result.movies["released"]]
    table = pd.DataFrame(
        {
            "Movie": upcoming["title"],
            "Release Date": upcoming["release_date"],
            "Predicted": upcoming["predicted"],
            "Tracker Projection": upcoming["projection"],
        }
    )
    _write_title(ws, "Upcoming: matched movies still waiting on an actual result", span_cols=4)
    header_row = 3
    last_row = _write_dataframe(
        ws, _for_excel(table), header_row, number_formats={1: "yyyy-mm-dd", 2: MONEY_FORMAT, 3: MONEY_FORMAT}
    )
    ws.column_dimensions["A"].width = 38
    _add_table_polish(ws, header_row, max(last_row, header_row), last_col=len(table.columns))


def _build_unmatched_sheet(ws, result: Comparison) -> None:
    _write_title(ws, "Unmatched: titles found in only one sheet", span_cols=2)
    length = max(len(result.unmatched_predictions), len(result.unmatched_actuals))
    table = pd.DataFrame(
        {
            "In predictions only": result.unmatched_predictions + [None] * (length - len(result.unmatched_predictions)),
            "In results only": result.unmatched_actuals + [None] * (length - len(result.unmatched_actuals)),
        }
    )
    header_row = 3
    last_row = _write_dataframe(ws, table, header_row)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 44
    _add_table_polish(ws, header_row, max(last_row, header_row), last_col=2)


def build_prediction_workbook(result: Comparison, findings: list[str], out_path: Path, as_of: date) -> None:
    wb = Workbook()
    scorecard_ws = wb.active
    scorecard_ws.title = "Scorecard"
    _build_scorecard_sheet(scorecard_ws, result, findings, as_of)
    _build_movies_sheet(wb.create_sheet("Movie by Movie"), result)
    _build_upcoming_sheet(wb.create_sheet("Upcoming"), result)
    _build_unmatched_sheet(wb.create_sheet("Unmatched"), result)
    for name, color in {"Scorecard": "1F4E78", "Movie by Movie": "2E7D32", "Upcoming": "6A1B9A", "Unmatched": "757575"}.items():
        wb[name].sheet_properties.tabColor = color
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def tableau_rows(result: Comparison) -> pd.DataFrame:
    """Long, tidy table for Tableau: one row per movie per forecaster that made a forecast."""
    frames = []
    for forecast, label in FORECASTS.items():
        movies = result.movies[result.movies[forecast].notna()]
        pct = movies[f"{forecast}_pct_error"]
        frames.append(
            pd.DataFrame(
                {
                    "movie": movies["title"],
                    "release_date": movies["release_date"].dt.strftime("%Y-%m-%d"),
                    "forecaster": label,
                    "forecast": movies[forecast],
                    "actual": movies["actual"].where(movies["released"]),
                    "status": movies["released"].map({True: "Released", False: "Upcoming"}),
                    "error": (movies[forecast] - movies["actual"]).where(movies["released"]),
                    "pct_error": pct.round(2),
                    "abs_pct_error": pct.abs().round(2),
                    "direction": pct.map(lambda v: "" if pd.isna(v) else "Too high" if v > 0 else "Too low" if v < 0 else "Exact"),
                    "accuracy_band": pct.abs().map(lambda v: "" if pd.isna(v) else accuracy_band(v)),
                }
            )
        )
    return pd.concat(frames, ignore_index=True)


def summary_markdown(result: Comparison, findings: list[str], as_of: date) -> str:
    lines = [f"## Prediction Accuracy as of {as_of}", ""]
    rows = scorecard(result)
    if rows:
        lines += ["| Measure | Result | Movies |", "|---|---|---|"]
        for row in rows:
            movies = ", ".join(row.movies).replace("|", "\\|")
            lines.append(f"| **{row.measure}** | {row.result} | {movies} |")
    lines += ["", "### Key findings", ""] + [f"- {finding}" for finding in findings]

    stats, headers = detailed_stats_rows(result)
    if headers:
        lines += ["", "<details><summary>Detailed stats</summary>", ""]
        lines += ["| Metric | " + " | ".join(headers) + " |", "|---|" + "--:|" * len(headers)]
        for label, values, kind in stats:
            lines.append(f"| {label} | " + " | ".join(format_value(v, kind) for v in values) + " |")
        lines += ["", "</details>"]
    return "\n".join(lines) + "\n"


def format_value(value: float, kind: str) -> str:
    if kind == "money":
        return short_money(value)
    if kind == "pct":
        return f"{value:+.1f}%"
    if kind == "pct_abs":
        return f"{value:.1f}%"
    return f"{value:,.0f}"


def run(
    predictions_source: str,
    actuals_source: str,
    as_of: date,
    title_map: list[tuple[str, str]] | None = None,
    history_path: Path = HISTORY_DB,
    save_history: bool = False,
) -> tuple[Comparison, list[str], Path]:
    """Scores the predictions and writes the report. This run's actuals and scores are added to an
    in-memory copy of the results history (blank actuals are filled from it); `save_history`
    writes that copy back to the history file."""
    log = logging.getLogger(__name__)
    actuals = load_actuals(actuals_source)
    predictions = read_predictions(load_source(predictions_source))

    with closing(open_history(history_path)) as conn:
        save_actuals(actuals, conn, as_of)
        actuals, restored = fill_from_history(actuals, conn)
        result = compare(predictions, actuals, title_map)
        result.restored_from_history = [t for t in restored if t in set(result.movies["title"])]
        scores = tableau_rows(result)
        save_scores(scores, conn, as_of)
        if save_history:
            write_history(conn, history_path)
            log.info("Saved this run's actuals and %d scores to %s", len(scores), history_path)

    findings = prediction_findings(result)
    out_path = next_free_path(OUTPUT_DIR / f"prediction_accuracy_{as_of}.xlsx")
    build_prediction_workbook(result, findings, out_path, as_of)
    scores.to_csv(out_path.with_suffix(".csv"), index=False)
    log.info("Prediction report written to %s (+ .csv for Tableau)", out_path)
    return result, findings, out_path


def _print_summary(result: Comparison, findings: list[str], out_path: Path, as_of: date) -> None:
    rows = scorecard(result)
    if rows:
        table = Table(title=f"Prediction Accuracy as of {as_of}", title_style="bold cyan", show_lines=True)
        table.add_column("Measure", style="bold")
        table.add_column("Result")
        table.add_column("Movies", ratio=1)
        for row in rows:
            table.add_row(escape(row.measure), escape(row.result), escape(", ".join(row.movies)))
        console.print()
        console.print(table)
    console.print("\n[bold cyan]Key findings[/bold cyan]")
    for finding in findings:
        console.print(f"  • {escape(finding)}")
    console.print(f"\n[bold green]Report written[/bold green] to [bold]{out_path}[/bold]")
    console.print(f"[bold green]Tableau-ready data[/bold green] in [bold]{out_path.with_suffix('.csv')}[/bold]\n")


def _parse_title_map(value: str) -> list[tuple[str, str]]:
    try:
        return parse_title_map(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError(str(exc)) from exc


def main() -> None:
    parser = argparse.ArgumentParser(description="Score movie box-office predictions against actual results.")
    parser.add_argument("--predictions", required=True, help="Link or file path for the predictions sheet.")
    parser.add_argument("--actuals", required=True, help="Link or file path for the actual-results sheet.")
    parser.add_argument(
        "--title-map",
        default=None,
        type=_parse_title_map,
        help='Optional: pair movies listed under different names, e.g. "Minions 3=MINIONS AND MONSTERS; '
        'Jumanji 3=JUMANJI: OPEN WORLD" (results title=predictions title, separated by semicolons).',
    )
    parser.add_argument(
        "--save-history",
        action="store_true",
        help="Also record this run's actual results in the results history database (used by the scheduled run).",
    )
    parser.add_argument("--markdown-summary", default=None, help="Optional: append the scorecard and findings as Markdown to this file.")
    args = parser.parse_args()

    as_of = date.today()
    setup_logging(as_of, "predictions")
    try:
        result, findings, out_path = run(
            args.predictions, args.actuals, as_of, args.title_map, save_history=args.save_history
        )
        if args.markdown_summary:
            with open(args.markdown_summary, "a", encoding="utf-8") as summary_file:
                summary_file.write(summary_markdown(result, findings, as_of))
    except Exception as exc:
        logging.getLogger(__name__).exception("Prediction report failed")
        console.print(f"[bold red]Error:[/bold red] {describe_error(exc)}")
        sys.exit(1)
    _print_summary(result, findings, out_path, as_of)


if __name__ == "__main__":
    main()
