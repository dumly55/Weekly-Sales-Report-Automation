"""Shared Excel styling for both reports: fonts, fills, and helpers for titles, tables and findings."""

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True, color="1F4E78")
SECTION_FONT = Font(size=12, bold=True, color="1F4E78")
BOLD = Font(bold=True)
ZEBRA_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
GOOD_FONT = Font(color="1F7A1F", bold=True)
BAD_FONT = Font(color="C00000", bold=True)

_GRID_SIDE = Side(style="thin", color="BFBFBF")
GRID_BORDER = Border(left=_GRID_SIDE, right=_GRID_SIDE, top=_GRID_SIDE, bottom=_GRID_SIDE)


def write_title(ws: Worksheet, text: str, span_cols: int, row: int = 1) -> None:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    if span_cols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span_cols)


def write_header_row(ws: Worksheet, headers: list[str], row: int, start_col: int = 1) -> None:
    for offset, header in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = GRID_BORDER


def write_dataframe(
    ws: Worksheet,
    df: pd.DataFrame,
    start_row: int,
    start_col: int = 1,
    number_formats: dict[int, str] | None = None,
    zebra: bool = True,
) -> int:
    """Writes a DataFrame with a styled, bordered header row and optional zebra striping.

    `number_formats` maps a 0-indexed column offset to a number format string,
    applied to every data cell in that column. Returns the last row written.
    """
    number_formats = number_formats or {}
    write_header_row(ws, list(df.columns), start_row, start_col)

    for r, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        is_alt_row = zebra and (r - start_row) % 2 == 0
        for c, value in enumerate(row):
            cell = ws.cell(row=r, column=start_col + c, value=value)
            cell.border = GRID_BORDER
            if c in number_formats:
                cell.number_format = number_formats[c]
            if is_alt_row:
                cell.fill = ZEBRA_FILL

    last_row = start_row + len(df)
    for offset, col in enumerate(df.columns):
        content_width = max((len(str(v)) for v in df[col]), default=0)
        width = min(40, max(12, len(str(col)) + 2, content_width + 2))
        ws.column_dimensions[get_column_letter(start_col + offset)].width = width

    return last_row


def add_table_polish(ws: Worksheet, header_row: int, last_row: int, last_col: int) -> None:
    """Adds an autofilter and freezes the header row so it stays visible while scrolling."""
    last_col_letter = get_column_letter(last_col)
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_row}"
    ws.freeze_panes = f"A{header_row + 1}"


FINDINGS_SPAN_COLS = 4
FINDINGS_CHARS_PER_LINE = 70


def write_findings(ws: Worksheet, findings: list[str], start_row: int) -> int:
    """Writes a "Key findings" block, one wrapped finding per row across the KPI table's width.
    Returns the last row used (start_row - 1 if there are no findings)."""
    if not findings:
        return start_row - 1
    ws.cell(row=start_row, column=1, value="Key findings").font = SECTION_FONT
    for offset, finding in enumerate(findings, start=1):
        row = start_row + offset
        cell = ws.cell(row=row, column=1, value=f"• {finding}")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=FINDINGS_SPAN_COLS)
        # Excel doesn't auto-fit the height of merged cells, so estimate it from the text length.
        lines = -(-len(finding) // FINDINGS_CHARS_PER_LINE)
        ws.row_dimensions[row].height = 15 * lines + 2
    return start_row + len(findings)
