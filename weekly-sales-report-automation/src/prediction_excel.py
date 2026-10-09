"""Builds the prediction accuracy Excel workbook."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .excel_style import (
    BAD_FONT,
    GOOD_FONT,
    GRID_BORDER,
    SECTION_FONT,
    add_table_polish,
    write_dataframe,
    write_findings,
    write_header_row,
    write_title,
)
from .prediction_findings import detailed_stats_rows, scorecard
from .predictions import AWAITING_RESULT, IN_THEATERS, UPCOMING, Comparison
from .sql_insights import Insight

GOOD_ERROR_PCT = 25  # errors within this are shown in green, beyond it in red
MONEY_FORMAT = '"$"#,##0'
PCT_FORMAT = "+0.0%;-0.0%"


MOVIES_COLUMN_WIDTH = 95


def _build_scorecard_sheet(ws, result: Comparison, findings: list[str], as_of: date) -> None:
    write_title(ws, f"Prediction Accuracy as of {as_of}", span_cols=3)
    rows = scorecard(result)
    if not rows:
        ws.cell(row=3, column=1, value="No matched movies to show yet.")
        write_findings(ws, findings, start_row=5)
        return

    header_row = 3
    write_header_row(ws, ["Measure", "Result", "Movies"], header_row)
    for offset, row in enumerate(rows, start=1):
        r = header_row + offset
        movies = ", ".join(row.movies)
        for col, value in enumerate([row.measure, row.result, movies], start=1):
            cell = ws.cell(row=r, column=col, value=value or None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = GRID_BORDER
        ws.cell(row=r, column=1).font = Font(bold=True)
        # Excel doesn't auto-fit wrapped rows written this way, so size each from its movie list.
        lines = max(1, -(-len(movies) // (MOVIES_COLUMN_WIDTH - 5)))
        ws.row_dimensions[r].height = 15 * lines + 2
    last_row = header_row + len(rows)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = MOVIES_COLUMN_WIDTH
    ws.freeze_panes = f"A{header_row + 1}"

    last_row = write_findings(ws, findings, start_row=last_row + 2)
    _write_detailed_stats(ws, result, start_row=last_row + 2)


def _write_detailed_stats(ws, result: Comparison, start_row: int) -> None:
    rows, headers = detailed_stats_rows(result)
    if not headers:  # nothing scored yet, so there are no stats to show
        return
    ws.cell(row=start_row, column=1, value="Detailed stats").font = SECTION_FONT
    table = pd.DataFrame([[label, *values] for label, values, _ in rows], columns=["Metric", *headers])
    header_row = start_row + 1
    write_dataframe(ws, table, header_row)
    formats = {"count": "#,##0", "money": MONEY_FORMAT, "pct": PCT_FORMAT, "pct_abs": "0.0%"}
    for offset, (_, values, kind) in enumerate(rows, start=1):
        for col in range(2, 2 + len(values)):
            cell = ws.cell(row=header_row + offset, column=col)
            if kind.startswith("pct"):
                cell.value = cell.value / 100
            cell.number_format = formats[kind]
    for col, width in {"A": 32, "B": 26, "C": MOVIES_COLUMN_WIDTH}.items():
        ws.column_dimensions[col].width = width


def _pct_cells(ws, header_row: int, last_row: int, columns: list[int]) -> None:
    for col in columns:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=col, max_col=col):
            if isinstance(cell.value, (int, float)):
                cell.font = GOOD_FONT if abs(cell.value) <= GOOD_ERROR_PCT else BAD_FONT
                cell.value = cell.value / 100
                cell.number_format = PCT_FORMAT


def _for_excel(df: pd.DataFrame) -> pd.DataFrame:
    return df.astype(object).where(df.notna(), None)


def _build_movies_sheet(ws, result: Comparison) -> None:
    released = result.movies[result.movies["scored"]].copy()

    def closer(row):
        p, t = row["predicted_pct_error"], row["projection_pct_error"]
        if pd.isna(p) or pd.isna(t):
            return None
        return "Predictions" if abs(p) < abs(t) else ("Tracker" if abs(t) < abs(p) else "Tie")

    table = pd.DataFrame(
        {
            "Movie": released["title"],
            "Release Date": released["release_date"],
            "Predicted": released["predicted"],
            "Tracker Projection": released["projection"],
            "Actual": released["actual"],
            "Predicted Error": released["predicted_pct_error"],
            "Tracker Error": released["projection_pct_error"],
            "Closer": released.apply(closer, axis=1) if len(released) else [],
        }
    )
    write_title(ws, "Movie by Movie: finished runs, forecasts vs final worldwide gross", span_cols=8)
    header_row = 3
    last_row = write_dataframe(
        ws, _for_excel(table), header_row, number_formats={1: "yyyy-mm-dd", 2: MONEY_FORMAT, 3: MONEY_FORMAT, 4: MONEY_FORMAT}
    )
    _pct_cells(ws, header_row, last_row, columns=[6, 7])
    ws.column_dimensions["A"].width = 38
    add_table_polish(ws, header_row, max(last_row, header_row), last_col=len(table.columns))


def _build_in_theaters_sheet(ws, result: Comparison) -> None:
    showing = result.movies[result.movies["run_status"] == IN_THEATERS]
    reached = showing["actual"] / showing["predicted"]
    table = pd.DataFrame(
        {
            "Movie": showing["title"],
            "Release Date": showing["release_date"],
            "Days in Theaters": showing["run_days"],
            "Gross So Far": showing["actual"],
            "Predicted": showing["predicted"],
            "% of Prediction Reached": reached,
            "Tracker Projection": showing["projection"],
            "% of Tracker Reached": showing["actual"] / showing["projection"],
            "Already Passed Prediction": (showing["actual"] > showing["predicted"]).map({True: "Yes", False: ""}),
        }
    ).sort_values("% of Prediction Reached", ascending=False, na_position="last")
    write_title(ws, "In Theaters: still earning, so not scored until the run ends", span_cols=9)
    header_row = 3
    last_row = write_dataframe(
        ws,
        _for_excel(table),
        header_row,
        number_formats={1: "yyyy-mm-dd", 3: MONEY_FORMAT, 4: MONEY_FORMAT, 5: "0%", 6: MONEY_FORMAT, 7: "0%"},
    )
    ws.column_dimensions["A"].width = 38
    add_table_polish(ws, header_row, max(last_row, header_row), last_col=len(table.columns))


def _build_upcoming_sheet(ws, result: Comparison) -> None:
    waiting = result.movies[result.movies["run_status"].isin([UPCOMING, AWAITING_RESULT])]
    waiting = waiting.sort_values(["run_status", "run_days", "release_date"], ascending=[False, True, True], na_position="last")
    table = pd.DataFrame(
        {
            "Movie": waiting["title"],
            "Status": waiting["run_status"].map(
                {UPCOMING: "Upcoming", AWAITING_RESULT: "Finished, no result in the sheet yet"}
            ),
            "Release Date": waiting["release_date"],
            "Days Until Release": waiting["run_days"].where(waiting["run_status"] == UPCOMING),
            "Predicted": waiting["predicted"],
            "Tracker Projection": waiting["projection"],
        }
    )
    write_title(ws, "Upcoming: not released yet, plus finished movies still missing a result", span_cols=6)
    header_row = 3
    last_row = write_dataframe(
        ws, _for_excel(table), header_row, number_formats={2: "yyyy-mm-dd", 4: MONEY_FORMAT, 5: MONEY_FORMAT}
    )
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 34
    add_table_polish(ws, header_row, max(last_row, header_row), last_col=len(table.columns))


def _build_unmatched_sheet(ws, result: Comparison) -> None:
    write_title(ws, "Unmatched: titles found in only one sheet", span_cols=2)
    length = max(len(result.unmatched_predictions), len(result.unmatched_actuals))
    table = pd.DataFrame(
        {
            "In predictions only": result.unmatched_predictions + [None] * (length - len(result.unmatched_predictions)),
            "In results only": result.unmatched_actuals + [None] * (length - len(result.unmatched_actuals)),
        }
    )
    header_row = 3
    last_row = write_dataframe(ws, table, header_row)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 44
    add_table_polish(ws, header_row, max(last_row, header_row), last_col=2)


def _build_insights_sheet(ws, insights: list[Insight]) -> None:
    write_title(ws, "SQL Insights: the queries in the sql/ folder, run on the results history", span_cols=6)
    row = 3
    for insight in insights:
        ws.cell(row=row, column=1, value=f"{insight.question}  ({insight.name}.sql)").font = SECTION_FONT
        table = insight.display_table
        if table.empty:
            ws.cell(row=row + 1, column=1, value="No rows yet.")
            row += 3
            continue
        row = write_dataframe(ws, _for_excel(table), row + 1) + 3
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 22


def build_prediction_workbook(
    result: Comparison, findings: list[str], out_path: Path, as_of: date, insights: list[Insight] | None = None
) -> None:
    wb = Workbook()
    scorecard_ws = wb.active
    scorecard_ws.title = "Scorecard"
    _build_scorecard_sheet(scorecard_ws, result, findings, as_of)
    _build_movies_sheet(wb.create_sheet("Movie by Movie"), result)
    _build_in_theaters_sheet(wb.create_sheet("In Theaters"), result)
    _build_upcoming_sheet(wb.create_sheet("Upcoming"), result)
    _build_unmatched_sheet(wb.create_sheet("Unmatched"), result)
    tab_colors = {
        "Scorecard": "1F4E78",
        "Movie by Movie": "2E7D32",
        "In Theaters": "C62828",
        "Upcoming": "6A1B9A",
        "Unmatched": "757575",
    }
    if insights:
        _build_insights_sheet(wb.create_sheet("SQL Insights"), insights)
        tab_colors["SQL Insights"] = "EF6C00"
    for name, color in tab_colors.items():
        wb[name].sheet_properties.tabColor = color
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
