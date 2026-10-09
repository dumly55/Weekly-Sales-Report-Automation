from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.prediction_report import (
    accuracy_band,
    build_prediction_workbook,
    prediction_findings,
    run,
    scorecard,
    short_money,
    summary_markdown,
    tableau_rows,
)
from src.predictions import compare
from src.sql_insights import Insight

AS_OF = date(2026, 10, 8)

COMPLETE = "✅ Run Complete"


def _result(with_results=True):
    """Finished runs: Alpha +10%, Bravo -50%, Charlie +300% (the tracker is exact on all three).
    Golf and Hotel are still in theaters (Hotel has already passed its prediction), Delta is upcoming,
    India finished but has no result in the sheet, Echo has no result row and Foxtrot no prediction."""
    predictions = pd.DataFrame(
        {
            "title": ["Alpha", "Bravo", "Charlie", "Delta", "Echo", "Golf", "Hotel", "India"],
            "predicted": [110.0, 50.0, 400.0, 100.0, 70.0, 200.0, 50.0, 30.0],
        }
    )
    actuals = pd.DataFrame(
        {
            "title": ["Alpha", "Bravo", "Charlie", "Delta", "Foxtrot", "Golf", "Hotel", "India"],
            "actual": [100.0, 100.0, 100.0, None, 90.0, 120.0, 80.0, None] if with_results else [None] * 8,
            "projection": [100.0, 100.0, 100.0, 120.0, 95.0, 150.0, 60.0, 40.0],
            "release_date": pd.to_datetime(
                ["2026-01-02", "2026-01-09", "2026-01-16", "2026-12-18", "2026-02-01", "2026-09-28", "2026-09-18", "2026-03-01"]
            ),
            "status_text": [
                COMPLETE, COMPLETE, COMPLETE, "70 Days", COMPLETE,
                "🎬 In Theaters for 10 days", "🎬 In Theaters for 20 days", COMPLETE,
            ],
        }
    )
    return compare(predictions, actuals)


class TestFindings:
    def test_findings_cover_accuracy_progress_and_unmatched(self):
        assert prediction_findings(_result()) == [
            "3 of the 7 matched movies have finished their run and are scored. Not scored yet: 2 still in theaters, "
            "1 not released yet and 1 finished but missing a result.",
            "The predictions were typically off by 50% (median; the average is 120% because of a few big misses). "
            "1 of 3 landed within 25% of the actual result, and 1 within 10%.",
            "They ran high more often than not: 2 too high vs 1 too low (median error +10%).",
            "Added up, the predictions came to $560 against $300 actual (+86.7%).",
            "Best call: Alpha, predicted $110 vs $100 actual (+10.0%).",
            "Biggest miss: Charlie, predicted $400 vs $100 actual (+300.0%).",
            "Against the tracker's own projections (typically off by 0%), the predictions were closer on 0 of 3 movies.",
            "Still in theaters, so not scored until their run ends: Hotel at $80 after 20 days (160% of its $50 "
            "prediction); Golf at $120 after 10 days (60% of its $200 prediction).",
            "Hotel has already earned more than predicted, so it was predicted too low whatever happens next.",
            "Next release: Delta opens in 70 days, predicted at $100.",
            "1 movie has finished its run but has no result in the results sheet yet, so it isn't scored: India.",
            "1 movie in the results sheet has no matching prediction (it may be listed under a different title): Foxtrot.",
            "1 predicted movie doesn't appear in the results sheet; see the Unmatched sheet.",
        ]

    def test_nothing_finished_yet(self):
        findings = prediction_findings(_result(with_results=False))
        assert findings[0] == (
            "0 of the 7 matched movies have finished their run and are scored. Not scored yet: 2 still in theaters, "
            "1 not released yet and 4 finished but missing a result."
        )
        assert not any("typically off" in f for f in findings)
        assert any("Golf (after 10 days, no gross reported yet)" in f for f in findings)


def test_scorecard_lists_the_movies_behind_each_measure():
    rows = [(r.measure, r.result, r.movies) for r in scorecard(_result())]
    assert rows == [
        ("Finished and scored", "3 of 7 matched movies", []),
        ("Typical miss (median)", "50%", []),
        ("Nailed it (within 10%)", "1 of 3", ["Alpha (+10%)"]),
        ("Close (10-25% off)", "0 of 3", []),
        ("Off (25-50% off)", "1 of 3", ["Bravo (-50%)"]),
        ("Way off (more than 50% off)", "1 of 3", ["Charlie (+300%)"]),
        ("Predicted too high", "2 of 3", ["Charlie (+300%)", "Alpha (+10%)"]),
        ("Predicted too low", "1 of 3", ["Bravo (-50%)"]),
        ("Closer than the tracker", "0 of 3", []),
        ("Tracker's typical miss (median)", "0%", []),
        ("Total predicted vs actual", "$560 vs $300 (+86.7%)", []),
        (
            "Still in theaters (not scored yet)",
            "2 of 7",
            ["Hotel ($80 after 20 days, 160% of prediction)", "Golf ($120 after 10 days, 60% of prediction)"],
        ),
        ("Upcoming", "1 of 7", ["Delta (in 70 days)"]),
        ("Finished, but no result in the sheet yet", "1 of 7", ["India"]),
    ]


def test_scorecard_before_any_results_shows_only_whats_not_scored():
    assert [r.measure for r in scorecard(_result(with_results=False))] == [
        "Still in theaters (not scored yet)",
        "Upcoming",
        "Finished, but no result in the sheet yet",
    ]


def test_accuracy_bands_have_inclusive_upper_limits():
    assert [accuracy_band(v) for v in (0, 10, 10.1, 25, 50, 50.1)] == ["Nailed it", "Nailed it", "Close", "Close", "Off", "Way off"]


def test_short_money():
    assert short_money(2_506_720_000) == "$2.51B"
    assert short_money(58_527_279) == "$58.5M"
    assert short_money(940_000) == "$940,000"


class TestWorkbook:
    def test_sheets_and_contents(self, tmp_path):
        result = _result()
        out_path = tmp_path / "accuracy.xlsx"
        build_prediction_workbook(result, prediction_findings(result), out_path, AS_OF)

        wb = load_workbook(out_path)
        assert wb.sheetnames == ["Scorecard", "Movie by Movie", "In Theaters", "Upcoming", "Unmatched"]

        ws = wb["Scorecard"]
        assert [c.value for c in ws[3]][:3] == ["Measure", "Result", "Movies"]
        assert [ws["A6"].value, ws["B6"].value, ws["C6"].value] == ["Nailed it (within 10%)", "1 of 3", "Alpha (+10%)"]
        detailed = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Detailed stats")
        assert ws.cell(row=detailed + 1, column=1).value == "Metric"
        assert ws.cell(row=detailed + 3, column=1).value == "Median error (either direction)"
        assert ws.cell(row=detailed + 3, column=2).value == pytest.approx(0.5)
        assert ws.cell(row=detailed + 3, column=2).number_format == "0.0%"

        movies = wb["Movie by Movie"]
        assert [movies.cell(row=r, column=1).value for r in (4, 5, 6, 7)] == ["Alpha", "Bravo", "Charlie", None]
        assert movies["F4"].value == pytest.approx(0.10)
        assert movies["H4"].value == "Tracker"

        showing = wb["In Theaters"]
        assert [showing["A4"].value, showing["C4"].value, showing["D4"].value] == ["Hotel", 20, 80]
        assert showing["F4"].value == pytest.approx(1.6) and showing["F4"].number_format == "0%"
        assert showing["I4"].value == "Yes"
        assert showing["A5"].value == "Golf" and showing["I5"].value is None

        upcoming = wb["Upcoming"]
        assert [upcoming["A4"].value, upcoming["B4"].value, upcoming["D4"].value] == ["Delta", "Upcoming", 70]
        assert [upcoming["A5"].value, upcoming["B5"].value] == ["India", "Finished, no result in the sheet yet"]
        assert [c.value for c in wb["Unmatched"][4]] == ["Echo", "Foxtrot"]

    def test_scorecard_before_any_results(self, tmp_path):
        result = _result(with_results=False)
        out_path = tmp_path / "accuracy.xlsx"
        build_prediction_workbook(result, prediction_findings(result), out_path, AS_OF)
        assert load_workbook(out_path)["Scorecard"]["A4"].value == "Still in theaters (not scored yet)"


def test_tableau_rows_are_one_per_movie_per_forecaster():
    rows = tableau_rows(_result())
    assert len(rows) == 14

    def row(movie, forecaster="Predictions"):
        return rows[(rows["movie"] == movie) & (rows["forecaster"] == forecaster)].iloc[0]

    alpha = row("Alpha")
    assert (alpha["release_date"], alpha["forecast"], alpha["actual"]) == ("2026-01-02", 110.0, 100.0)
    assert (alpha["status"], alpha["error"], alpha["pct_error"], alpha["direction"]) == ("Final", 10.0, 10.0, "Too high")
    assert alpha["accuracy_band"] == "Nailed it" and pd.isna(alpha["gross_so_far"])

    golf = row("Golf")
    assert golf["status"] == "In theaters"
    assert (golf["gross_so_far"], golf["pct_of_forecast_reached"], golf["days_in_theaters"]) == (120.0, 60.0, 10)
    assert pd.isna(golf["actual"]) and pd.isna(golf["pct_error"]) and golf["accuracy_band"] == ""

    delta = row("Delta", "Tracker Projection")
    assert (delta["status"], delta["days_until_release"]) == ("Upcoming", 70)
    assert pd.isna(delta["pct_error"]) and delta["direction"] == ""


def test_run_end_to_end_with_local_files(tmp_path, monkeypatch):
    monkeypatch.setattr("src.prediction_report.OUTPUT_DIR", tmp_path)
    (tmp_path / "predictions.csv").write_text(
        "MOVIE TITLE,WORLDWIDE TOTAL\nMERCY,\"$41,470,588.24\"\nDIGGER,$349000000\n", encoding="utf-8"
    )
    (tmp_path / "actuals.csv").write_text(
        "Report,,,\nMovie,Release Date,Countdown,Worldwide Actual\n"
        "Mercy 👮🏻,23/01/2026,✅ Run Complete,\"$54,709,856\"\n"
        "Digger,02/10/2026,🎬 In Theaters for 6 days,\"$24,947,690\"\n",
        encoding="utf-8",
    )
    output = run(
        str(tmp_path / "predictions.csv"), str(tmp_path / "actuals.csv"), AS_OF, history_path=tmp_path / "h.sqlite"
    )

    assert output.out_path == tmp_path / "prediction_accuracy_2026-10-08.xlsx" and output.out_path.exists()
    assert len(pd.read_csv(output.out_path.with_suffix(".csv"))) == 2
    assert output.findings[0] == (
        "1 of the 2 matched movies have finished their run and are scored. Not scored yet: 1 still in theaters."
    )
    assert any("Digger at $24.9M after 6 days (7% of its $349.0M prediction)" in f for f in output.findings)

    assert len(output.insights) == 8
    assert "SQL Insights" in load_workbook(output.out_path).sheetnames
    history = pd.read_csv(output.history_csv)
    assert list(history["status"]) == ["In theaters", "Final"]


def test_summary_markdown():
    result = _result()
    markdown = summary_markdown(result, prediction_findings(result), AS_OF)
    assert markdown.startswith("## Prediction Accuracy as of 2026-10-08\n")
    assert "| **Nailed it (within 10%)** | 1 of 3 | Alpha (+10%) |" in markdown
    assert "| **Upcoming** | 1 of 7 | Delta (in 70 days) |" in markdown
    assert "<details><summary>Detailed stats</summary>" in markdown
    assert "| Median error (either direction) | 50.0% | 0.0% |" in markdown
    assert "- Best call: Alpha, predicted $110 vs $100 actual (+10.0%)." in markdown
    assert "SQL insights" not in markdown


def test_summary_markdown_with_sql_insights():
    result = _result()
    insights = [
        Insight("02_median_miss", "What is the median miss?", pd.DataFrame({"forecaster": ["Predictions"], "median_miss_pct": [50.0]})),
        Insight("03_accuracy_bands", "Which band?", pd.DataFrame({"accuracy_band": ["Off"], "movies": [1], "titles": ["Bravo"]})),
    ]
    markdown = summary_markdown(result, prediction_findings(result), AS_OF, insights)
    assert "<details><summary>SQL insights (queries in the sql/ folder)</summary>" in markdown
    assert "**What is the median miss?** (`02_median_miss.sql`)" in markdown
    assert "| Predictions | 50.0 |" in markdown
    assert "| accuracy_band | movies |" in markdown, "the long titles column is left out"
